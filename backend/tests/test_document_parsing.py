import asyncio
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

import app.services.document_service as document_service_module
from app.services.document_service import DocumentService, SavedDocumentFile


class FakeUploadFile:
    def __init__(
        self,
        filename: str,
        content: bytes,
        content_type: str = "application/octet-stream",
    ):
        self.filename = filename
        self.content_type = content_type
        self._file = BytesIO(content)

    async def read(self, size: int = -1) -> bytes:
        return self._file.read(size)


def create_service(tmp_path: Path) -> DocumentService:
    return DocumentService(
        upload_dir=tmp_path / "uploads",
        parsed_dir=tmp_path / "parsed",
    )


def save_file(
    service: DocumentService,
    filename: str,
    content: bytes,
) -> SavedDocumentFile:
    return asyncio.run(service.save_upload_file(FakeUploadFile(filename, content)))


def parse_file(service: DocumentService, saved_file: SavedDocumentFile):
    return asyncio.run(service.parse_saved_file(saved_file))


def assert_no_parsed_files(service: DocumentService) -> None:
    if service.parsed_dir.exists():
        assert not any(path.is_file() for path in service.parsed_dir.rglob("*"))


def test_parse_txt_and_save_metadata(tmp_path):
    service = create_service(tmp_path)
    saved_file = save_file(service, "notes.txt", b"hello\nworld")

    result = parse_file(service, saved_file)
    parsed_path = service.build_storage_paths(saved_file.id, ".txt").parsed_file

    assert result.id == saved_file.id
    assert result.display_filename == "notes.txt"
    assert result.extension == ".txt"
    assert result.text == "hello\nworld"
    assert result.character_count == len("hello\nworld")
    assert parsed_path == service.parsed_dir / saved_file.id / "content.txt"
    assert parsed_path.read_text(encoding="utf-8") == result.text


def test_parse_txt_with_utf8_bom(tmp_path):
    service = create_service(tmp_path)
    saved_file = save_file(service, "bom.txt", b"\xef\xbb\xbfhello")

    result = parse_file(service, saved_file)

    assert result.text == "hello"


def test_parse_txt_with_invalid_utf8_is_rejected(tmp_path):
    service = create_service(tmp_path)
    saved_file = save_file(service, "invalid.txt", b"\xff\xfeinvalid")

    with pytest.raises(ValueError, match="UTF-8 encoded"):
        parse_file(service, saved_file)

    assert service.build_storage_paths(saved_file.id, ".txt").upload_file.exists()
    assert_no_parsed_files(service)


def test_parse_pdf_joins_non_empty_pages(tmp_path, monkeypatch):
    service = create_service(tmp_path)
    saved_file = save_file(service, "guide.pdf", b"fake pdf content")
    pages = [
        SimpleNamespace(extract_text=lambda: "First page"),
        SimpleNamespace(extract_text=lambda: None),
        SimpleNamespace(extract_text=lambda: "Second page"),
    ]
    monkeypatch.setattr(
        document_service_module,
        "PdfReader",
        lambda _: SimpleNamespace(pages=pages),
    )

    result = parse_file(service, saved_file)

    assert result.text == "First page\n\nSecond page"


def test_parse_pdf_without_text_is_rejected(tmp_path):
    service = create_service(tmp_path)
    pdf_content = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(pdf_content)
    saved_file = save_file(service, "blank.pdf", pdf_content.getvalue())

    with pytest.raises(ValueError, match="no parseable text"):
        parse_file(service, saved_file)

    assert service.build_storage_paths(saved_file.id, ".pdf").upload_file.exists()
    assert_no_parsed_files(service)


def test_parse_docx_extracts_paragraphs_and_table_cells(tmp_path):
    service = create_service(tmp_path)
    document = DocxDocument()
    document.add_paragraph("Introduction")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "AnythingLLM Mini"
    document.add_paragraph("Conclusion")
    docx_content = BytesIO()
    document.save(docx_content)
    saved_file = save_file(service, "guide.docx", docx_content.getvalue())

    result = parse_file(service, saved_file)

    assert result.text == "Introduction\nName\tAnythingLLM Mini\nConclusion"


def test_parse_whitespace_only_document_is_rejected(tmp_path):
    service = create_service(tmp_path)
    saved_file = save_file(service, "empty.txt", b" \n\t ")

    with pytest.raises(ValueError, match="no parseable text"):
        parse_file(service, saved_file)

    assert service.build_storage_paths(saved_file.id, ".txt").upload_file.exists()
    assert_no_parsed_files(service)


def test_parse_corrupt_document_is_rejected(tmp_path):
    service = create_service(tmp_path)
    saved_file = save_file(service, "broken.docx", b"not a docx archive")

    with pytest.raises(ValueError, match="failed to parse document"):
        parse_file(service, saved_file)

    assert service.build_storage_paths(saved_file.id, ".docx").upload_file.exists()
    assert_no_parsed_files(service)


def test_parse_rejects_source_symlink_outside_upload_directory(tmp_path):
    service = create_service(tmp_path)
    outside_path = tmp_path / "outside.txt"
    outside_path.write_text("outside", encoding="utf-8")
    saved_file = SavedDocumentFile(
        id="a" * 32,
        display_filename="outside.txt",
        content_type="text/plain",
        size_bytes=7,
        extension=".txt",
    )
    upload_path = service.build_storage_paths(saved_file.id, ".txt").upload_file
    upload_path.parent.mkdir(parents=True)
    upload_path.symlink_to(outside_path)

    with pytest.raises(ValueError, match="invalid upload path"):
        parse_file(service, saved_file)

    assert outside_path.exists()
    assert_no_parsed_files(service)


def test_parse_rejects_extension_mismatch(tmp_path):
    service = create_service(tmp_path)
    saved_file = save_file(service, "notes.txt", b"notes")
    mismatched_file = saved_file.model_copy(update={"extension": ".pdf"})

    with pytest.raises(ValueError, match="extension does not match"):
        parse_file(service, mismatched_file)

    assert service.build_storage_paths(saved_file.id, ".txt").upload_file.exists()
    assert_no_parsed_files(service)
